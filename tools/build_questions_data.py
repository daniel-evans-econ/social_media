"""Build the committed question bank from the answer-key files.

This is a one-off / occasional DEV script (not imported at runtime). It:
  1. Reads the answer keys in ``social_media/questions/<task>/``.
  2. Copies + normalizes the question images into ``social_media/static/questions/<task>/``
     so oTree can serve them.
  3. Writes ``social_media/questions_data.py`` (QUESTIONS, SETS, ...), which the
     app imports at runtime (so the app never needs openpyxl / Excel).

Re-run after adding new images (e.g. the real shape-rotation stimuli):

    python tools/build_questions_data.py

Requires the dev dependency ``openpyxl`` (see requirements-dev.txt).
"""
from __future__ import annotations

import csv
import os
import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
QUESTIONS_DIR = ROOT / "social_media" / "questions"
STATIC_DIR = ROOT / "social_media" / "static" / "questions"
OUT_FILE = ROOT / "social_media" / "questions_data.py"

# Difficulty ordering used everywhere.
DIFFICULTIES = ["easy", "medium", "hard"]

# Multiple-choice option labels per task.
TASK_OPTIONS = {
    "sequences": ["A", "B", "C", "D"],
    "shape_rotation": ["A", "B", "C", "D", "E"],
    "ravens": ["1", "2", "3", "4", "5", "6", "7", "8"],
}

# How a task is answered: "mc" = single multiple-choice; "count" = type the
# number of red-dot squares (working memory).
TASK_RESPONSE = {
    "sequences": "mc",
    "shape_rotation": "mc",
    "ravens": "mc",
    "working_memory": "count",
}

WM_GRID_ROWS = 3
WM_GRID_COLS = 10


def _rows_from_xlsx(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        out.append(dict(zip(header, r)))
    return out


def _copy_image(src: Path, task: str, item_id: str) -> str:
    """Copy ``src`` into the static dir as ``<item_id><ext>``; return the static-relative path."""
    ext = src.suffix.lower()
    dest_dir = STATIC_DIR / task
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{item_id}{ext}"
    shutil.copyfile(src, dest)
    return f"questions/{task}/{item_id}{ext}"


def build_sequences():
    src_dir = QUESTIONS_DIR / "sequences"
    rows = _rows_from_xlsx(src_dir / "sequences_answer_key.xlsx")
    # Map item_id -> source image. Filenames: question_<NN>_<diff>_<NN>.png
    images = {}
    for f in src_dir.glob("*.png"):
        tokens = f.stem.split("_")  # ['question','01','easy','01']
        item_id = f"{tokens[-2]}_{tokens[-1]}"
        images[item_id] = f
    items = {}
    for r in rows:
        item_id = str(r["Item ID"]).strip()
        difficulty = str(r["Difficulty"]).strip().lower()
        correct = str(r["Correct Option"]).strip()
        image = _copy_image(images[item_id], "sequences", item_id) if item_id in images else None
        items[item_id] = dict(
            item_id=item_id, difficulty=difficulty, image=image,
            options=list(TASK_OPTIONS["sequences"]), correct=correct,
            is_placeholder=image is None,
        )
    return items


def build_shape_rotation():
    src_dir = QUESTIONS_DIR / "shape_rotation"
    rows = _rows_from_xlsx(src_dir / "shape_rotation_answer_key.xlsx")
    # The 30 items are ordered by increasing difficulty, so split into
    # easy (1-10) / medium (11-20) / hard (21-30) to match the other tasks
    # (two fixed sets of five per difficulty).
    # Images are named <item_id>.png (e.g. 3d_shapes_07.png); the example.png
    # is the worked example and is intentionally excluded from the pool.
    images = {f.stem: f for f in src_dir.glob("*.png")}
    items = {}
    for r in rows:
        item_id = str(r["Item ID"]).strip()
        item_no = int(r["Item No."])
        correct = str(r["Correct Option"]).strip()
        if item_no <= 10:
            difficulty = "easy"
        elif item_no <= 20:
            difficulty = "medium"
        else:
            difficulty = "hard"
        src = images.get(item_id)
        image = _copy_image(src, "shape_rotation", item_id) if src else None
        items[item_id] = dict(
            item_id=item_id, difficulty=difficulty, image=image,
            options=list(TASK_OPTIONS["shape_rotation"]), correct=correct,
            is_placeholder=image is None,
        )
    return items


def build_ravens():
    src_dir = QUESTIONS_DIR / "ravens_matrices"
    rows = _rows_from_xlsx(src_dir / "ravens_answer_key.xlsx")
    # Images are named <item>.jpg (e.g. C1.jpg). Exactly the 10 items per block
    # that should be in the live bank are present as files; we drive inclusion
    # off the images present (the example item, e.g. C8, is intentionally absent).
    images = {}
    for f in src_dir.iterdir():
        if f.suffix.lower() in (".jpg", ".jpeg", ".png") and f.stem.lower() != "example":
            images[f.stem.upper()] = f
    block_to_diff = {"C": "easy", "D": "medium", "E": "hard"}
    items = {}
    for r in rows:
        block = str(r["Block"]).strip().upper()
        item = str(r["Item"]).strip().upper()
        if block not in block_to_diff:
            continue
        if item not in images:
            continue  # not provided -> excluded from the live bank
        image = _copy_image(images[item], "ravens", item)
        items[item] = dict(
            item_id=item, difficulty=block_to_diff[block], image=image,
            options=list(TASK_OPTIONS["ravens"]), correct=str(int(r["Correct answer"])),
            is_placeholder=False,
        )
    return items


def build_working_memory():
    src_dir = QUESTIONS_DIR / "working_memory"
    items = {}
    with open(src_dir / "working_memory_answer_key.csv", newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            fname = r["file"].strip()
            item_id = os.path.splitext(fname)[0]
            difficulty = r["difficulty"].strip().lower()
            occupied = [c for c in r["occupied_cells"].split() if c]
            src = src_dir / fname
            image = _copy_image(src, "working_memory", item_id) if src.exists() else None
            items[item_id] = dict(
                item_id=item_id, difficulty=difficulty, image=image,
                occupied_cells=occupied, dot_count=int(r["dot_count"]),
                grid_rows=WM_GRID_ROWS, grid_cols=WM_GRID_COLS,
                is_placeholder=image is None,
            )
    return items


def _sets_by_difficulty(items: dict) -> dict:
    """Two fixed sets of five per difficulty: <diff>_1, <diff>_2."""
    sets = {}
    for diff in DIFFICULTIES:
        ids = sorted(i for i, v in items.items() if v["difficulty"] == diff)
        for idx in range(0, len(ids), 5):
            chunk = ids[idx:idx + 5]
            if len(chunk) == 5:
                sets[f"{diff}_{idx // 5 + 1}"] = chunk
    return sets


def _sets_single_pool(items: dict) -> dict:
    ids = sorted(items.keys())
    sets = {}
    for idx in range(0, len(ids), 5):
        chunk = ids[idx:idx + 5]
        if len(chunk) == 5:
            sets[f"set_{idx // 5 + 1}"] = chunk
    return sets


def main():
    if STATIC_DIR.exists():
        shutil.rmtree(STATIC_DIR)
    sequences = build_sequences()
    shape_rotation = build_shape_rotation()
    ravens = build_ravens()
    working_memory = build_working_memory()

    questions = {
        "sequences": sequences,
        "shape_rotation": shape_rotation,
        "ravens": ravens,
        "working_memory": working_memory,
    }
    sets = {
        "sequences": _sets_by_difficulty(sequences),
        "shape_rotation": _sets_by_difficulty(shape_rotation),
        "ravens": _sets_by_difficulty(ravens),
        "working_memory": _sets_by_difficulty(working_memory),
    }

    import pprint
    header = (
        '"""AUTO-GENERATED by tools/build_questions_data.py - do not edit by hand.\n\n'
        'Run ``python tools/build_questions_data.py`` to regenerate after changing\n'
        'the answer keys or adding question images.\n"""\n\n'
    )
    body = (
        f"TASK_RESPONSE = {pprint.pformat(TASK_RESPONSE, sort_dicts=True)}\n\n"
        f"TASK_OPTIONS = {pprint.pformat(TASK_OPTIONS, sort_dicts=True)}\n\n"
        f"WM_GRID_ROWS = {WM_GRID_ROWS}\n"
        f"WM_GRID_COLS = {WM_GRID_COLS}\n\n"
        f"QUESTIONS = {pprint.pformat(questions, sort_dicts=True, width=120)}\n\n"
        f"SETS = {pprint.pformat(sets, sort_dicts=True, width=120)}\n"
    )
    OUT_FILE.write_text(header + body, encoding="utf-8")

    print(f"Wrote {OUT_FILE}")
    for task, its in questions.items():
        placeholders = sum(1 for v in its.values() if v["is_placeholder"])
        print(f"  {task}: {len(its)} items ({placeholders} placeholder), {len(sets[task])} sets")


if __name__ == "__main__":
    main()
